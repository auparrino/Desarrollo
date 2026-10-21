"""Parse INDEC Censo 2022 definitive provincial cuadros (hogares/vivienda/pob) to a
departamento-level table. Usage: python3 -I parse_indec.py <dir_cuadros> <out_csv>"""
import sys, re, glob, os, unicodedata
import openpyxl
import pandas as pd

D, OUT = sys.argv[1], sys.argv[2]

def num(x):
    if x is None: return None
    s = str(x).replace('\xa0', '').replace(' ', '').replace('.', '')
    if s in ('-', ''): return 0
    if s == '///': return None
    try: return int(float(s))
    except ValueError: return None

def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)

def code(x):
    if x is None: return None
    s = str(x).strip().replace('\xa0', '')
    if re.fullmatch(r'\d{4,5}', s): return s.zfill(5)
    return None

def rows(ws):
    return [[c for c in r] for r in ws.iter_rows(values_only=True)]

# Province files: c2022_<slug>_<tema>_c<n>_<k>.xlsx (k omitted for 'tp')
files = sorted(glob.glob(os.path.join(D, 'c2022_*_c*.xlsx')))
prov = {}
for f in files:
    m = re.match(r'c2022_\s*([a-z]+)_(hogares|vivienda|pob)_c(\d+)(?:_\d+)?\.xlsx', os.path.basename(f))
    if not m or m.group(1) == 'tp': continue
    prov.setdefault(m.group(1), {})[(m.group(2), int(m.group(3)))] = f

def dept_table(f, ncols):
    """Single 'según departamento' table: rows with a 5-digit code."""
    wb = openpyxl.load_workbook(f, read_only=True)
    out = {}
    for sn in wb.sheetnames:
        if not sn.replace('\xa0', '').startswith('Cuadro'): continue
        for r in rows(wb[sn]):
            c = code(r[0]) if r else None
            if c and c not in out:
                out[c] = [str(r[1]).replace('\xa0', ' ').strip()] + [num(v) for v in r[2:2 + ncols]]
    return out

def per_dept_sheets(f):
    """Per-department sheets 'Cuadro X.k.j': returns {name_norm: rows}."""
    wb = openpyxl.load_workbook(f, read_only=True)
    out = []
    for sn in wb.sheetnames:
        s = re.sub(r'\s', '', sn.replace('\xa0', ''))
        if not re.fullmatch(r'Cuadro\d+\.\d+\.\d+', s): continue
        rr = rows(wb[sn])
        title = str(rr[1][0]).replace('\xa0', ' ')
        m = re.search(r'(?:departamento|partido|comuna|Comuna)\s+(.+?)\.\s+(?:Total|Población)', title, re.I)
        name = m.group(1) if m else title
        if re.search(r'Buenos Aires,?\s*Comuna\s+\d+', title):
            name = 'Comuna ' + re.search(r'Comuna\s+(\d+)', title).group(1)
        out.append((name, rr))
    return out

def find_row(rr, label, col):
    for r in rr:
        if r and r[0] is not None and norm(r[0]).startswith(norm(label)):
            # c3 has an extra label column ('Total') before numbers
            vals = [v for v in r[1:] if v is not None and str(v).strip() != '' and not norm(v).startswith('total') or str(v).strip() in ('-',)]
            return r
    return None

recs = []
for slug, fs in sorted(prov.items()):
    gas = dept_table(fs[('hogares', 4)], 7)     # total, elec, gas red, tubo, garrafa, leña, otro
    ten = dept_table(fs[('hogares', 6)], 10)    # total, propia(total,escr,boleto,otra,sindoc), alquilada, cedida, prestada, otra
    viv = dept_table(fs[('vivienda', 1)], 10)   # tot viv, viv part, presentes, vacac, oficina, alq/venta, constr, ausentes, otra, colectivas
    vho = dept_table(fs[('vivienda', 2)], 2)    # viv part ocupadas, hogares
    pob = dept_table(fs[('pob', 4)], 1)         # poblacion en viviendas particulares
    name2code = {}
    for c, v in gas.items():
        if c[2:] == '000': continue
        name2code[norm(v[0])] = c
    def attach(fkey, label, key):
        res = {}
        for name, rr in per_dept_sheets(fs[fkey]):
            alias = {'nuevedejulio': '9dejulio', 'lezana': 'lezama',
                     'coroneldemarinalrosales': 'coroneldemarinaleonardorosales'}
            c = name2code.get(alias.get(norm(name), norm(name)))
            if c is None:
                cands = [k for n, k in name2code.items() if norm(name) in n or n in norm(name)]
                c = cands[0] if len(cands) == 1 else None
            # verify with the total of households (first 'Total' row); fix INDEC title errors
            tot = None
            for r in rr:
                if r and r[0] is not None and norm(r[0]) == 'total':
                    tot = next((num(v) for v in r[1:] if v is not None and str(v).strip() != '' and num(v) is not None), None); break
            if c is None or c in res or gas[c][1] != tot:
                cands = [k for k in gas if k[2:] != '000' and k not in res and gas[k][1] == tot]
                if len(cands) == 1:
                    print('INFO', slug, fkey, repr(name), '->', cands[0], gas[cands[0]][0], '(por total de hogares)', file=sys.stderr)
                    c = cands[0]
                else:
                    print('WARN no code for', slug, fkey, name, file=sys.stderr); continue
            res[c] = rr
        return res
    agua = attach(('hogares', 2), 'Red pública', 'agua')
    clo = attach(('hogares', 3), 'A red pública', 'cloaca')
    net = attach(('hogares', 7), 'Total', 'internet')
    for c, g in gas.items():
        if len(c) != 5 or c[2:] == '000' : continue
        rec = {'cod_dep': c, 'departamento_indec': g[0], 'prov_slug': slug}
        rec['hogares'] = g[1]
        # '///' (no corresponde) en gas de red = departamento sin red: se registra 0
        rec['hog_gas_red_cocina'] = g[3] if g[3] is not None else 0
        rec['hog_electricidad_cocina'] = g[2]
        rec['hog_garrafa_cocina'] = g[5]
        rec['hog_lena_carbon_cocina'] = g[6]
        t = ten.get(c)
        if t:
            rec.update(hog_tot_ten=t[1], hog_propia=t[2], hog_propia_escritura=t[3], hog_alquilada=t[7],
                       hog_cedida_trabajo=t[8], hog_prestada=t[9], hog_otra_tenencia=t[10])
        v = viv.get(c)
        if v:
            rec.update(viv_total=v[1], viv_part=v[2], viv_part_con_personas=v[3], viv_vacaciones=v[4],
                       viv_oficina=v[5], viv_alquiler_venta=v[6], viv_construccion=v[7],
                       viv_moradores_ausentes=v[8], viv_otra_situacion=v[9], viv_colectivas=v[10])
        vh = vho.get(c)
        if vh: rec.update(viv_part_ocupadas=vh[1], hogares_v2=vh[2])
        p = pob.get(c)
        if p: rec['pob_viv_part'] = p[1]
        rr = agua.get(c)
        if rr:
            for r in rr:
                if r and r[0] and norm(r[0]).startswith('total') and 'hog_tot_agua' not in rec: rec['hog_tot_agua'] = num(r[1])
                if r and r[0] and norm(r[0]).startswith('redpublica'): rec['hog_agua_red'] = num(r[1])
        rr = clo.get(c)
        if rr:
            for r in rr:
                if r and r[0] and norm(r[0]).startswith('total') and 'hog_tot_desague' not in rec: rec['hog_tot_desague'] = num(r[2])
                if r and r[0] and norm(r[0]).startswith('aredpublica'): rec['hog_cloaca'] = num(r[2])
        rr = net.get(c)
        if rr:
            for r in rr:
                if r and r[0] and norm(r[0]) == 'total' and num(r[1]) is not None:
                    rec['hog_tot_internet'] = num(r[1]); rec['hog_internet'] = num(r[2]); break
        recs.append(rec)

df = pd.DataFrame(recs).sort_values('cod_dep')
df.to_csv(OUT, index=False)
print(len(df), 'departamentos'); print(df.isna().sum()[df.isna().sum() > 0])
