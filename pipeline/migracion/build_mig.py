"""Build web/src/model/migracion.json from MIALC (CEPAL/CELADE) tables + census series.
Usage: python3 -I build_mig.py <raw_dir> <scratch_dir> <out_json> <stats_out>
"""
import sys, json, math, datetime as dt
import openpyxl
import numpy as np

RAW, SCR, OUT, STATS = sys.argv[1:5]

PROV = ["Ciudad Autónoma de Buenos Aires", "Buenos Aires", "Catamarca", "Córdoba", "Corrientes", "Chaco",
        "Chubut", "Entre Ríos", "Formosa", "Jujuy", "La Pampa", "La Rioja", "Mendoza", "Misiones", "Neuquén",
        "Río Negro", "Salta", "San Juan", "San Luis", "Santa Cruz", "Santa Fe", "Santiago del Estero", "Tucumán",
        "Tierra del Fuego, Antártida e Islas del Atlántico Sur"]

def n(v):
    if v is None: return 0
    s = str(v).replace(" ", "").replace("\xa0", "")
    return 0 if s in ("", "-") else int(round(float(s)))

# ---------- 2022 DAM matrix (rows = residence 2022 = destination; cols = residence 2017 = origin)
ws = openpyxl.load_workbook(f"{RAW}/m2022/ARG_2022_D5MB.xlsx", data_only=True, read_only=True).worksheets[0]
rows = [list(r) for r in ws.iter_rows(values_only=True)]
hdr = rows[13]
assert [h.strip() if isinstance(h, str) else h for h in hdr[2:26]] == PROV, hdr[2:26]
M22 = {}  # M22[dest][orig]
for k in range(24):
    r = rows[14 + k]
    M22[PROV[k]] = {PROV[j]: n(r[2 + j]) for j in range(24)}
    assert n(r[26]) == sum(M22[PROV[k]].values())
tot_row = rows[38]
assert str(tot_row[1]).strip() == "Total"
mat22 = {o: {d: M22[d][o] for d in PROV} for o in PROV}   # mat22[orig][dest]
inter22 = sum(mat22[o][d] for o in PROV for d in PROV if o != d)
assert inter22 == 750405, inter22

# ---------- 2010 DAM matrix
ws = openpyxl.load_workbook(f"{RAW}/m2010/ARG_2010_D5MB.xlsx", data_only=True, read_only=True).worksheets[0]
rows10 = [list(r) for r in ws.iter_rows(values_only=True)]
M10 = {}
for k in range(24):
    r = rows10[4 + k]
    M10[PROV[k]] = {PROV[j]: n(r[2 + j]) for j in range(24)}
mat10 = {o: {d: M10[d][o] for d in PROV} for o in PROV}
inter10 = sum(mat10[o][d] for o in PROV for d in PROV if o != d)
assert inter10 == 1224216, inter10

# ---------- 2010 AMBA split (aggregated from DAME matrix in agg3.py)
a10 = json.load(open(f"{SCR}/mig_2010_amba.json"))
order = a10["order"]; R = a10["R"]   # R[orig][dest]
# consistency check: aggregating AMBA+Resto back to "Buenos Aires"+CABA should reproduce DAM flows of other provinces
for p in PROV[2:]:
    assert R[p][p] == mat10[p][p] or abs(R[p][p] - mat10[p][p]) < 20, (p, R[p][p], mat10[p][p])
inter10amba = sum(R[o][d] for o in order for d in order if o != d)

def summary(mat, regs):
    out = {}
    for r in regs:
        inm = sum(mat[o][r] for o in regs if o != r)
        emi = sum(mat[r][d] for d in regs if d != r)
        pop_dest = sum(mat[o][r] for o in regs)
        pop_orig = sum(mat[r][d] for d in regs)
        out[r] = {"inmigrantes": inm, "emigrantes": emi, "saldo": inm - emi,
                  "poblacion_5mas_fin": pop_dest, "poblacion_5mas_inicio": pop_orig,
                  "tasa_saldo_pct_anio": round(100 * (inm - emi) / 5 / ((pop_dest + pop_orig) / 2), 3),
                  "tasa_inmig_pct_anio": round(100 * inm / 5 / ((pop_dest + pop_orig) / 2), 3)}
    return out

s22 = summary(mat22, PROV); s10 = summary(mat10, PROV); s10a = summary(R, order)

def corridors(mat, regs, k=25):
    c = [(o, d, mat[o][d]) for o in regs for d in regs if o != d]
    c.sort(key=lambda x: -x[2])
    return c[:k]

def origin_shares(mat, regs, dest, k=6):
    tot = sum(mat[o][dest] for o in regs if o != dest)
    l = sorted(((o, mat[o][dest]) for o in regs if o != dest), key=lambda x: -x[1])[:k]
    return [(o, v, round(100 * v / tot, 1)) for o, v in l]

# ---------- sex by destination 2022 (immigrants), from D5SE
ws = openpyxl.load_workbook(f"{RAW}/m2022/ARG_2022_D5SE.xlsx", data_only=True, read_only=True).worksheets[0]
rs = [list(r) for r in ws.iter_rows(values_only=True)]
masc = {}
k = -1
cur = None
for r in rs[14:]:
    if r[1] not in (None, "") and str(r[1]).strip() == "Total":
        break
    if r[1] not in (None, ""):
        k += 1
        if k >= 24: break
        cur = PROV[k]; fem = mas = None
    lab = str(r[2] or "")
    if cur and "Mujer" in lab:
        fem = sum(n(r[3 + j]) for j in range(24) if j != PROV.index(cur))
    if cur and "Var" in lab:
        mas = sum(n(r[3 + j]) for j in range(24) if j != PROV.index(cur))
        masc[cur] = round(100 * mas / fem, 1) if fem else None
        assert fem + mas == s22[cur]["inmigrantes"], (cur, fem, mas, s22[cur]["inmigrantes"])

# ---------- distance decay fit (2022), capitals/main city coordinates
COORD = {  # main urban centre of each jurisdiction (lat, lon)
    "Ciudad Autónoma de Buenos Aires": (-34.61, -58.38), "Buenos Aires": (-34.75, -58.60),
    "Catamarca": (-28.47, -65.78), "Córdoba": (-31.42, -64.18), "Corrientes": (-27.47, -58.83),
    "Chaco": (-27.45, -58.99), "Chubut": (-45.86, -67.48), "Entre Ríos": (-31.73, -60.53),
    "Formosa": (-26.18, -58.18), "Jujuy": (-24.19, -65.30), "La Pampa": (-36.62, -64.29),
    "La Rioja": (-29.41, -66.86), "Mendoza": (-32.89, -68.83), "Misiones": (-27.37, -55.90),
    "Neuquén": (-38.95, -68.06), "Río Negro": (-39.03, -67.58), "Salta": (-24.79, -65.41),
    "San Juan": (-31.54, -68.54), "San Luis": (-33.30, -66.34), "Santa Cruz": (-51.62, -69.22),
    "Santa Fe": (-32.95, -60.64), "Santiago del Estero": (-27.80, -64.26), "Tucumán": (-26.82, -65.22),
    "Tierra del Fuego, Antártida e Islas del Atlántico Sur": (-53.79, -67.71)}
def hav(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (*a, *b))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))
def fit(mat, regs, popkey):
    X, y = [], []
    for o in regs:
        for d in regs:
            if o == d: continue
            if {o, d} == {"Ciudad Autónoma de Buenos Aires", "Buenos Aires"}: continue
            v = mat[o][d]
            if v <= 0: continue
            D = hav(COORD[o], COORD[d])
            X.append([1, math.log(popkey[o]), math.log(popkey[d]), math.log(D)]); y.append(math.log(v))
    X = np.array(X); y = np.array(y)
    b, *_ = np.linalg.lstsq(X, y, rcond=None)
    yhat = X @ b; r2 = 1 - ((y - yhat) ** 2).sum() / ((y - y.mean()) ** 2).sum()
    return {"const": round(b[0], 3), "elast_pob_origen": round(b[1], 3), "elast_pob_destino": round(b[2], 3),
            "elast_distancia": round(b[3], 3), "r2_log": round(r2, 3), "n_pares": len(y)}
pop22 = {p: s22[p]["poblacion_5mas_inicio"] for p in PROV}
pop10 = {p: s10[p]["poblacion_5mas_inicio"] for p in PROV}
fit22 = fit(mat22, PROV, pop22); fit10 = fit(mat10, PROV, pop10)
# share of interprovincial migrants going to an adjacent-ish destination (< 500 km between main centres)
def share_dist(mat, regs, km):
    t = s = 0
    for o in regs:
        for d in regs:
            if o == d: continue
            v = mat[o][d]; t += v
            if hav(COORD[o], COORD[d]) < km: s += v
    return round(100 * s / t, 1)
dist_shares = {f"<{km}km": share_dist(mat22, PROV, km) for km in (300, 500, 1000)}
# distance bands (mean migrants per pair / (Po*Pd)) to show decay
def bands(mat, regs, pop):
    edges = [0, 300, 600, 1000, 1500, 2000, 4000]; acc = {}
    for o in regs:
        for d in regs:
            if o == d or {o, d} == {"Ciudad Autónoma de Buenos Aires", "Buenos Aires"}: continue
            D = hav(COORD[o], COORD[d])
            for a, b in zip(edges, edges[1:]):
                if a <= D < b:
                    k = f"{a}-{b}"; x = acc.setdefault(k, [0, 0.0])
                    x[0] += mat[o][d]; x[1] += pop[o] * pop[d] / 1e12
    return {k: round(v[0] / v[1], 1) for k, v in acc.items()}
bands22 = bands(mat22, PROV, pop22)

# ---------- profile
prof = json.load(open(f"{SCR}/mig_perfil.json"))
def typ(path, year):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    for r in ws.iter_rows(values_only=True):
        if r[1] and str(r[1]).strip() == "Total":
            v = [n(x) for x in r[2:8]]
            return {"directos_antiguos": v[0], "directos_recientes": v[1], "multiples": v[2], "retorno": v[3], "no_migrantes": v[4], "total": v[5]}
t22 = typ(f"{RAW}/m2022/ARG_2022_DTSX.xlsx", 2022); t10 = typ(f"{RAW}/m2010/ARG_2010_DTSX.xlsx", 2010)
def rec_shares(t):
    rec = t["directos_recientes"] + t["multiples"] + t["retorno"]
    return {"recientes_total": rec, "retorno_pct": round(100 * t["retorno"] / rec, 1),
            "multiples_pct": round(100 * t["multiples"] / rec, 1), "directos_pct": round(100 * t["directos_recientes"] / rec, 1)}
ts22 = rec_shares(t22); ts10 = rec_shares(t10)

# family indicator from age table
def fam(path, mig_first):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    import re
    ages = {}
    for r in ws.iter_rows(values_only=True):
        a = r[1]
        if a and str(a).strip()[:1].isdigit():
            lo = int(re.match(r"\s*(\d+)", str(a)).group(1))
            m, nm = (n(r[2]), n(r[3])) if mig_first else (n(r[3]), n(r[2]))
            ages[lo] = (m, nm)
    ch = lambda i: sum(v[i] for a, v in ages.items() if 5 <= a <= 14)
    ad = lambda i: sum(v[i] for a, v in ages.items() if 25 <= a <= 44)
    return round(ch(0) / ad(0), 3), round(ch(1) / ad(1), 3)
fam22 = fam(f"{RAW}/m2022/ARG_2022_D5DA.xlsx", False); fam10 = fam(f"{RAW}/m2010/ARG_2010_D5DA.xlsx", True)

sex22 = (368841, 381564); sex22nm = (19895458, 21490744)
sex10 = (612263, 611953); sex10nm = (16787519, 17958442)
e22 = prof["2022"]; e10 = prof["2010"]

# ---------- pace benchmarks
D = {"1970": dt.date(1970, 9, 30), "1980": dt.date(1980, 10, 22), "1991": dt.date(1991, 5, 15),
     "2001": dt.date(2001, 11, 17), "2010": dt.date(2010, 10, 27), "2022": dt.date(2022, 5, 18)}
def yrs(a, b): return (D[b] - D[a]).days / 365.25
def g(p0, p1, a, b): return round(100 * ((p1 / p0) ** (1 / yrs(a, b)) - 1), 2)
def absg(p0, p1, a, b): return round((p1 - p0) / yrs(a, b))

ritmos = []
def add(lugar, periodo, net_anio, crec, fuente, **kw):
    e = {"lugar": lugar, "periodo": periodo, "migrantes_netos_anio": net_anio, "crec_pct_anio": crec, "fuente": fuente}
    e.update(kw); ritmos.append(e)

CP = "citypop"; TDF = "tdf_doctec15"; MI22 = "mialc_2022"; MI10 = "mialc_2010_dam"; MI10M = "mialc_2010_dame"; MI01 = "mialc_2001"
RNP = "renaper_interprov_2023"; RND = "renaper_interdep_2023"; NQ = "formosa_cuadro_hist"
# --- Tierra del Fuego (industrial promotion, Ley 19.640 de 1972)
add("Tierra del Fuego (provincia)", "1970-1980", None, g(13527, 27358, "1970", "1980"), TDF,
    poblacion_inicio=13527, poblacion_fin=27358, crec_abs_anio=absg(13527, 27358, "1970", "1980"),
    notas="Crecimiento total (natural + migratorio). Promoción industrial Ley 19.640 (1972). Saldo migratorio no publicado: null.")
add("Tierra del Fuego (provincia)", "1980-1991", None, g(27358, 69369, "1980", "1991"), TDF,
    poblacion_inicio=27358, poblacion_fin=69369, crec_abs_anio=absg(27358, 69369, "1980", "1991"),
    notas="Período de crecimiento más rápido registrado para una provincia argentina; total, no sólo migración.")
add("Tierra del Fuego (provincia)", "1996-2001", round(3122 / 5), g(69369, 101079, "1991", "2001"), MI01,
    tasa_migr_neta_pct_anio=round(100 * 3122 / 5 / ((89317 + 86195) / 2), 2),
    notas="Saldo interprovincial 5 años (MIALC 2001); crec_pct_anio es el intercensal 1991-2001 (citypopulation/INDEC).")
add("Tierra del Fuego (provincia)", "2005-2010", round(s10["Tierra del Fuego, Antártida e Islas del Atlántico Sur"]["saldo"] / 5),
    g(101079, 127205, "2001", "2010"), MI10, tasa_migr_neta_pct_anio=s10["Tierra del Fuego, Antártida e Islas del Atlántico Sur"]["tasa_saldo_pct_anio"],
    notas="crec_pct_anio intercensal 2001-2010.")
add("Tierra del Fuego (provincia)", "2017-2022", round(s22["Tierra del Fuego, Antártida e Islas del Atlántico Sur"]["saldo"] / 5),
    g(127205, 185732, "2010", "2022"), MI22, tasa_migr_neta_pct_anio=s22["Tierra del Fuego, Antártida e Islas del Atlántico Sur"]["tasa_saldo_pct_anio"],
    notas="Saldo casi nulo pese a la rotación más alta del país (inmigración 1,56% y emigración 1,52% de la población por año). crec_pct_anio intercensal 2010-2022.")
add("Río Grande + Tolhuin (deptos., TdF)", "1991-2001", None, g(39371 + 445, 53930 + 1201, "1991", "2001"), CP,
    poblacion_inicio=39816, poblacion_fin=55131, crec_abs_anio=absg(39816, 55131, "1991", "2001"))
add("Río Grande + Tolhuin (deptos., TdF)", "2010-2022", None, g(67093 + 2949, 99241 + 6039, "2010", "2022"), CP,
    poblacion_inicio=70042, poblacion_fin=105280, crec_abs_anio=absg(70042, 105280, "2010", "2022"),
    notas="Tolhuin se separó de Río Grande; se suman para comparar.")
add("Ushuaia (depto., TdF)", "1991-2001", None, g(29411, 45785, "1991", "2001"), CP,
    poblacion_inicio=29411, poblacion_fin=45785, crec_abs_anio=absg(29411, 45785, "1991", "2001"))
add("Ushuaia (depto., TdF)", "2010-2022", None, g(56956, 80371, "2010", "2022"), CP,
    poblacion_inicio=56956, poblacion_fin=80371, crec_abs_anio=absg(56956, 80371, "2010", "2022"))
# --- Neuquén / oil
add("Neuquén (provincia)", "1970-1980", None, g(154570, 243850, "1970", "1980"), NQ,
    poblacion_inicio=154570, poblacion_fin=243850, crec_abs_anio=absg(154570, 243850, "1970", "1980"),
    notas="Hidroeléctricas (El Chocón), petróleo y Estado provincial; total, no sólo migración.")
add("Neuquén (provincia)", "1980-1991", None, g(243850, 388833, "1980", "1991"), NQ,
    poblacion_inicio=243850, poblacion_fin=388833, crec_abs_anio=absg(243850, 388833, "1980", "1991"))
add("Neuquén (provincia)", "2005-2010", round(s10["Neuquén"]["saldo"] / 5), g(474155, 551266, "2001", "2010"), MI10,
    tasa_migr_neta_pct_anio=s10["Neuquén"]["tasa_saldo_pct_anio"])
add("Neuquén (provincia)", "2012-2022", round(24484 / 10.6), g(551266, 710814, "2010", "2022"), RNP,
    notas="Saldo RENAPER (cambios de domicilio en DNI) ene-2012 a ago-2022 = 24.484; dividido por 10,6 años. Subregistra mudanzas no declaradas. crec_pct_anio intercensal 2010-2022.")
add("Neuquén (provincia)", "2017-2022", round(s22["Neuquén"]["saldo"] / 5), g(551266, 710814, "2010", "2022"), MI22,
    tasa_migr_neta_pct_anio=s22["Neuquén"]["tasa_saldo_pct_anio"])
add("Confluencia (depto., Neuquén)", "1991-2001", None, g(265123, 314793, "1991", "2001"), CP,
    poblacion_inicio=265123, poblacion_fin=314793, crec_abs_anio=absg(265123, 314793, "1991", "2001"))
add("Confluencia (depto., Neuquén)", "2010-2022", round(17035 / 10.6), g(362673, 468794, "2010", "2022"), CP,
    poblacion_inicio=362673, poblacion_fin=468794, crec_abs_anio=absg(362673, 468794, "2010", "2022"),
    notas="migrantes_netos_anio: saldo interdepartamental RENAPER 2012-2022 (17.035) / 10,6 años (fuente renaper_interdep_2023).")
add("Añelo (depto., Neuquén; Vaca Muerta)", "2010-2022", None, g(10786, 18166, "2010", "2022"), CP,
    poblacion_inicio=10786, poblacion_fin=18166, crec_abs_anio=absg(10786, 18166, "2010", "2022"),
    notas="RENAPER: 391 inmigrantes interdepartamentales por mil hab. en 2012-2022 (7ª tasa del país). El intendente sostiene que el censo subestima (>10.000 hab. en la localidad).")
add("Añelo (localidad)", "2010-2022", None, g(2449, 6304, "2010", "2022"), CP,
    poblacion_inicio=2449, poblacion_fin=6304, crec_abs_anio=absg(2449, 6304, "2010", "2022"))
add("Rincón de los Sauces (localidad, Neuquén; petróleo)", "1991-2001", None, g(3475, 10071, "1991", "2001"), CP,
    poblacion_inicio=3475, poblacion_fin=10071, crec_abs_anio=absg(3475, 10071, "1991", "2001"))
add("Rincón de los Sauces (localidad, Neuquén; petróleo)", "2001-2010", None, g(10071, 18691, "2001", "2010"), CP,
    poblacion_inicio=10071, poblacion_fin=18691, crec_abs_anio=absg(10071, 18691, "2001", "2010"))
add("Plottier (localidad, Neuquén)", "2010-2022", None, g(32390, 52291, "2010", "2022"), CP,
    poblacion_inicio=32390, poblacion_fin=52291, crec_abs_anio=absg(32390, 52291, "2010", "2022"))
add("Plottier (municipio, Neuquén)", "2005-2010", round(2929 / 5), g(22874, 32390, "2001", "2010"), MI10M,
    tasa_migr_neta_pct_anio=round(100 * 2929 / 5 / 30243, 2),
    notas="Saldo intermunicipal (incluye mudanzas desde la ciudad de Neuquén). crec_pct_anio: localidad 2001-2010.")
# --- Chubut / Santa Cruz
add("Puerto Madryn (localidad, Chubut)", "2001-2010", None, g(57614, 81315, "2001", "2010"), CP,
    poblacion_inicio=57614, poblacion_fin=81315, crec_abs_anio=absg(57614, 81315, "2001", "2010"))
add("Puerto Madryn (municipio, Chubut)", "2005-2010", round(4264 / 5), None, MI10M,
    tasa_migr_neta_pct_anio=round(100 * 4264 / 5 / 70424, 2), notas="Saldo con todos los demás municipios del país.")
add("Puerto Madryn (localidad, Chubut)", "2010-2022", None, g(81315, 97625, "2010", "2022"), CP,
    poblacion_inicio=81315, poblacion_fin=97625, crec_abs_anio=absg(81315, 97625, "2010", "2022"))
add("Santa Cruz (provincia)", "2005-2010", round(s10["Santa Cruz"]["saldo"] / 5), g(196958, 273964, "2001", "2010"), MI10,
    tasa_migr_neta_pct_anio=s10["Santa Cruz"]["tasa_saldo_pct_anio"], notas="Boom de obra pública, petróleo y minería.")
add("Santa Cruz (provincia)", "2017-2022", round(s22["Santa Cruz"]["saldo"] / 5), g(273964, 337226, "2010", "2022"), MI22,
    tasa_migr_neta_pct_anio=s22["Santa Cruz"]["tasa_saldo_pct_anio"], notas="Reversión: de mayor saldo relativo del país a expulsora.")
add("Lago Argentino (depto., Santa Cruz; El Calafate, turismo)", "1991-2001", None, g(3940, 7500, "1991", "2001"), CP,
    poblacion_inicio=3940, poblacion_fin=7500, crec_abs_anio=absg(3940, 7500, "1991", "2001"))
add("Lago Argentino (depto., Santa Cruz; El Calafate, turismo)", "2001-2010", None, g(7500, 18864, "2001", "2010"), CP,
    poblacion_inicio=7500, poblacion_fin=18864, crec_abs_anio=absg(7500, 18864, "2001", "2010"),
    notas="Crecimiento sostenido más rápido entre los departamentos de más de 5.000 hab. revisados; destino turístico.")
add("Lago Argentino (depto., Santa Cruz; El Calafate, turismo)", "2010-2022", None, g(18864, 25586, "2010", "2022"), CP,
    poblacion_inicio=18864, poblacion_fin=25586, crec_abs_anio=absg(18864, 25586, "2010", "2022"))
# --- Andean amenity
add("San Carlos de Bariloche (localidad, Río Negro)", "2001-2010", None, g(91851, 112205, "2001", "2010"), CP,
    poblacion_inicio=91851, poblacion_fin=112205, crec_abs_anio=absg(91851, 112205, "2001", "2010"))
add("San Carlos de Bariloche (municipio, Río Negro)", "2005-2010", round(459 / 5), None, MI10M,
    tasa_migr_neta_pct_anio=round(100 * 459 / 5 / 96199, 2), notas="Alta rotación (8.738 inmigrantes, 8.279 emigrantes en 5 años) con saldo bajo.")
add("San Carlos de Bariloche (localidad, Río Negro)", "2010-2022", None, g(112205, 134978, "2010", "2022"), CP,
    poblacion_inicio=112205, poblacion_fin=134978, crec_abs_anio=absg(112205, 134978, "2010", "2022"))
# --- AMBA periphery
add("Pilar (partido, Buenos Aires)", "2001-2010", None, g(232463, 299077, "2001", "2010"), CP,
    poblacion_inicio=232463, poblacion_fin=299077, crec_abs_anio=absg(232463, 299077, "2001", "2010"))
add("Pilar (partido, Buenos Aires)", "2005-2010", round(14413 / 5), None, MI10M,
    tasa_migr_neta_pct_anio=round(100 * 14413 / 5 / 259511, 2), notas="Saldo intermunicipal. De 30.066 inmigrantes, 12.497 venían de otra provincia (sobre todo CABA) y el resto de otros partidos bonaerenses.")
add("Pilar (partido, Buenos Aires)", "2010-2022", round(26037 / 10.6), g(299077, 394754, "2010", "2022"), CP,
    poblacion_inicio=299077, poblacion_fin=394754, crec_abs_anio=absg(299077, 394754, "2010", "2022"),
    notas="migrantes_netos_anio: saldo interdepartamental RENAPER 2012-2022 (26.037, el mayor del país) / 10,6 años.")
add("Escobar (partido, Buenos Aires)", "1991-2001", None, g(128421, 178155, "1991", "2001"), CP,
    poblacion_inicio=128421, poblacion_fin=178155, crec_abs_anio=absg(128421, 178155, "1991", "2001"))
add("Escobar (partido, Buenos Aires)", "2010-2022", None, g(213619, 256268, "2010", "2022"), CP,
    poblacion_inicio=213619, poblacion_fin=256268, crec_abs_anio=absg(213619, 256268, "2010", "2022"))
add("General Rodríguez (partido, Buenos Aires)", "2010-2022", round(20830 / 10.6), g(87185, 142709, "2010", "2022"), CP,
    poblacion_inicio=87185, poblacion_fin=142709, crec_abs_anio=absg(87185, 142709, "2010", "2022"),
    notas="migrantes_netos_anio: saldo RENAPER 2012-2022 / 10,6.")
add("San Vicente (partido, Buenos Aires)", "2010-2022", round(9599 / 10.6), g(59478, 98215, "2010", "2022"), CP,
    poblacion_inicio=59478, poblacion_fin=98215, crec_abs_anio=absg(59478, 98215, "2010", "2022"),
    notas="migrantes_netos_anio: saldo RENAPER 2012-2022 / 10,6.")
# --- Córdoba
add("Córdoba Capital (depto.)", "2010-2022", round(-50584 / 10.6), g(1329604, 1505250, "2010", "2022"), CP,
    poblacion_inicio=1329604, poblacion_fin=1505250, crec_abs_anio=absg(1329604, 1505250, "2010", "2022"),
    notas="Saldo RENAPER negativo: la capital pierde hacia Colón, Santa María y Punilla (suburbanización).")
add("Colón (depto., Córdoba; periurbano de Córdoba)", "2001-2010", None, g(171067, 225151, "2001", "2010"), CP,
    poblacion_inicio=171067, poblacion_fin=225151, crec_abs_anio=absg(171067, 225151, "2001", "2010"))
add("Colón (depto., Córdoba; periurbano de Córdoba)", "2010-2022", round(23525 / 10.6), g(225151, 295725, "2010", "2022"), CP,
    poblacion_inicio=225151, poblacion_fin=295725, crec_abs_anio=absg(225151, 295725, "2010", "2022"),
    notas="migrantes_netos_anio: saldo RENAPER 2012-2022 / 10,6.")
add("Santa María (depto., Córdoba)", "2010-2022", round(19942 / 10.6), g(98188, 144829, "2010", "2022"), CP,
    poblacion_inicio=98188, poblacion_fin=144829, crec_abs_anio=absg(98188, 144829, "2010", "2022"),
    notas="migrantes_netos_anio: saldo RENAPER 2012-2022 / 10,6.")
add("Córdoba (provincia)", "2017-2022", round(s22["Córdoba"]["saldo"] / 5), g(3308876, 3840905, "2010", "2022"), MI22,
    tasa_migr_neta_pct_anio=s22["Córdoba"]["tasa_saldo_pct_anio"], notas="Mayor saldo interprovincial absoluto en 2017-2022.")
add("San Luis (provincia; promoción industrial)", "1996-2001", round(10513 / 5), None, MI01,
    tasa_migr_neta_pct_anio=round(100 * 10513 / 5 / ((326627 + 316114) / 2), 2),
    notas="Saldo interprovincial MIALC 2001.")
add("San Luis (provincia)", "2017-2022", round(s22["San Luis"]["saldo"] / 5), None, MI22,
    tasa_migr_neta_pct_anio=s22["San Luis"]["tasa_saldo_pct_anio"])
add("AMBA (CABA + 24 partidos)", "2005-2010", round(s10a["AMBA"]["saldo"] / 5), None, MI10M,
    tasa_migr_neta_pct_anio=s10a["AMBA"]["tasa_saldo_pct_anio"], notas="Saldo con el resto del país (incluye resto de Buenos Aires). Negativo.")

# ---------- annual totals
renaper_tot = 5637155  # interdepartamental, ene-2012 a ago-2022
renaper_interprov = sum([295288, 202589, 92582, 96411, 56254, 80692, 100327, 53048, 70125, 31988, 99229, 125844, 24287,
                         28023, 58474, 54318, 23850, 44874, 67875, 46163, 74063, 82710, 62152, 472452])

out = {
    "fuente": "CEPAL/CELADE, base de datos MIALC 2.0 (Migración Interna en América Latina y el Caribe), tabla ARG_2022_D5MB 'Matriz básica DAM, migración reciente', procesada con REDATAM sobre los microdatos del Censo Nacional de Población, Hogares y Viviendas 2022 (INDEC). Pregunta censal 17-18: provincia de residencia habitual hace 5 años.",
    "periodo": "2017-2022",
    "unidad": "personas en el período",
    "universo": "Población de 5 años o más en viviendas particulares, censada en 2022. Se excluyen ignorados y quienes vivían en el extranjero en 2017. Mide personas (sobrevivientes al censo), no movimientos: no cuenta mudanzas múltiples ni a quienes migraron y murieron o emigraron al exterior.",
    "orientacion": "matriz[origen][destino]. origen = provincia de residencia habitual en 2017; destino = provincia de residencia habitual en 2022. La diagonal (origen = destino) es la población que no cambió de provincia; excluirla al calcular flujos.",
    "nota_amba": "El Censo 2022 sólo pregunta la provincia de residencia anterior (no el partido), por lo que la matriz 2022 no permite separar AMBA de resto de Buenos Aires en el origen. La matriz 2005-2010 con AMBA separado está en 'matriz_2010_amba'.",
    "regiones": PROV,
    "matriz": mat22,
    "resumen_regiones": s22,
    "corredores_principales": [{"origen": o, "destino": d, "personas": v} for o, d, v in corridors(mat22, PROV, 30)],
    "matriz_2010_amba": {
        "fuente": "CEPAL/CELADE MIALC 2.0, tabla ARG_2010_E5MB (matriz DAME municipal, migración reciente, Censo 2010 cuestionario ampliado, INDEC), agregada a 25 regiones: AMBA = CABA + 24 partidos del GBA; resto de Buenos Aires = 108 partidos restantes. Agregación propia.",
        "periodo": "2005-2010",
        "unidad": "personas en el período",
        "universo": "Población de 5 años o más; excluye ignorados y residentes en el extranjero en 2005.",
        "orientacion": "matriz[origen][destino]; diagonal = no migrantes entre regiones (incluye mudanzas entre municipios dentro de la misma región).",
        "regiones": order,
        "matriz": R,
        "resumen_regiones": s10a,
        "total_migrantes_entre_regiones": inter10amba,
        "corredores_principales": [{"origen": o, "destino": d, "personas": v} for o, d, v in corridors(R, order, 30)],
    },
    "matriz_2010_provincias": {
        "fuente": "CEPAL/CELADE MIALC 2.0, tabla ARG_2010_D5MB (Censo 2010, INDEC).",
        "periodo": "2005-2010", "unidad": "personas en el período", "regiones": PROV, "matriz": mat10,
        "resumen_regiones": s10},
    "perfil": {
        "edad_15_34_pct": e22["edad"]["m_15_34"],
        "edad_15_34_pct_no_migrantes": e22["edad"]["nm_15_34"],
        "edad_20_34_pct": e22["edad"]["m_20_34"],
        "edad_5_14_pct": e22["edad"]["m_5_14"],
        "edad_35_59_pct": e22["edad"]["m_35_59"],
        "edad_60mas_pct": e22["edad"]["m_60mas"],
        "edad_60mas_pct_no_migrantes": e22["edad"]["nm_60mas"],
        "tasa_migracion_5anios_pct_total": e22["edad"]["tasa_migr_total_pct"],
        "tasa_migracion_5anios_pct_20_34": e22["edad"]["tasa_migr_20_34_pct"],
        "tasa_migracion_5anios_pct_por_edad": e22["edad"]["tasa_por_edad"],
        "anios_estudio_25_39_migrantes": e22["educ"]["25-39"]["m"]["anios_medios"],
        "anios_estudio_25_39_no_migrantes": e22["educ"]["25-39"]["nm"]["anios_medios"],
        "pct_13mas_anios_estudio_25_39_migrantes": e22["educ"]["25-39"]["m"]["pct_13mas"],
        "pct_13mas_anios_estudio_25_39_no_migrantes": e22["educ"]["25-39"]["nm"]["pct_13mas"],
        "indice_masculinidad_migrantes": round(100 * sex22[0] / sex22[1], 1),
        "indice_masculinidad_no_migrantes": round(100 * sex22nm[0] / sex22nm[1], 1),
        "indice_masculinidad_inmigrantes_por_destino": masc,
        "indice_masculinidad_renaper_2012_2022": 116,
        "edad_mediana_renaper_2012_2022": 31,
        "ninos_5_14_por_adulto_25_44_migrantes": fam22[0],
        "ninos_5_14_por_adulto_25_44_no_migrantes": fam22[1],
        "retorno_pct_de_recientes": ts22["retorno_pct"],
        "multiples_pct_de_recientes": ts22["multiples_pct"],
        "directos_pct_de_recientes": ts22["directos_pct"],
        "comparacion_2010": {
            "edad_15_34_pct": e10["edad"]["m_15_34"], "edad_5_14_pct": e10["edad"]["m_5_14"], "edad_60mas_pct": e10["edad"]["m_60mas"],
            "tasa_migracion_5anios_pct_total": e10["edad"]["tasa_migr_total_pct"],
            "tasa_migracion_5anios_pct_20_34": e10["edad"]["tasa_migr_20_34_pct"],
            "anios_estudio_25_39_migrantes": e10["educ"]["25-39"]["m"]["anios_medios"],
            "anios_estudio_25_39_no_migrantes": e10["educ"]["25-39"]["nm"]["anios_medios"],
            "indice_masculinidad_migrantes": round(100 * sex10[0] / sex10[1], 1),
            "indice_masculinidad_no_migrantes": round(100 * sex10nm[0] / sex10nm[1], 1),
            "ninos_5_14_por_adulto_25_44_migrantes": fam10[0], "ninos_5_14_por_adulto_25_44_no_migrantes": fam10[1],
            "retorno_pct_de_recientes": ts10["retorno_pct"], "multiples_pct_de_recientes": ts10["multiples_pct"]},
        "decaimiento_distancia": {
            "modelo": "log(M_od) = c + a·log(P_o) + b·log(P_d) − g·log(D_od), MCO sobre pares interprovinciales con flujo > 0, excluido el par CABA–Buenos Aires. D = distancia de gran círculo entre centros urbanos principales (coordenadas aproximadas). Estimación propia, orientativa.",
            "ajuste_2017_2022": fit22, "ajuste_2005_2010": fit10,
            "pct_migrantes_a_menos_de_km_2017_2022": dist_shares,
            "intensidad_por_banda_km_2017_2022": {"definicion": "migrantes por cada 1e12 de producto P_o·P_d, por banda de distancia", "valores": bands22}},
        "notas": "Fuente principal: MIALC 2.0 tablas ARG_2022_D5DA (edad), D5EE (edad × años de estudio), D5SX (sexo), D5SE (matriz por sexo), DTSX (tipología) y sus equivalentes 2010. 'Migrante' = cambió de provincia entre 2017 y 2022. Años de estudio: excluye código 99 (ignorado); la codificación de años de estudio de 2010 y 2022 difiere, comparar sólo dentro de cada censo. Tipología CELADE (personas con lugar de nacimiento y residencia anterior conocidos): retorno = vivía en otra provincia en 2017 y en 2022 volvió a la de nacimiento; múltiples = la provincia de 2017 no es ni la de nacimiento ni la actual. El indicador de 'familias' es un cociente de edades (niños 5-14 por adulto 25-44), no una medición directa de hogares que migran juntos: no hay tabulado de migración por hogar. RENAPER (cambios de domicilio en DNI, 2012-2022) da edad mediana 31 e índice de masculinidad 116; el censo da 96,7, de modo que la sobrerrepresentación masculina del registro probablemente refleje quién actualiza el DNI.",
    },
    "ritmos": ritmos,
    "total_interprovincial_anio": round(inter22 / 5),
    "total_interprovincial_detalle": {
        "censo_2022_mialc_2017_2022_personas": inter22, "censo_2022_por_anio": round(inter22 / 5),
        "censo_2010_mialc_2005_2010_personas": inter10, "censo_2010_por_anio": round(inter10 / 5),
        "censo_2001_mialc_1996_2001_personas": 1076836, "censo_2001_por_anio": round(1076836 / 5),
        "renaper_interprovincial_2012_2022_personas": renaper_interprov, "renaper_interprovincial_por_anio": round(renaper_interprov / 10.6),
        "renaper_interdepartamental_2012_2022_personas": renaper_tot, "renaper_interdepartamental_por_anio": round(renaper_tot / 10.6),
        "notas": "El dato censal de 5 años subestima los movimientos anuales (cuenta personas, no mudanzas; omite a quienes volvieron dentro del quinquenio o migraron siendo menores de 5). Por eso el registro RENAPER (con AMBA = CABA + 40 partidos y resto de Buenos Aires como jurisdicción aparte) da más migrantes por año. La caída 2010→2022 (de ~245 mil a ~150 mil por año) coincide con lo que reporta CEPAL ('fuerte caída' de la intensidad), pero puede estar exagerada por la pandemia de 2020-21 y por el cambio a censo de derecho y autoempadronamiento digital en 2022; usar 150-250 mil personas/año como rango."},
    "fuentes": [
        {"id": "mialc_2022", "cita": "CEPAL/CELADE (2025-2026). MIALC 2.0 – Migración Interna en América Latina y el Caribe. Argentina 2022, División Administrativa Mayor: tablas ARG_2022_D5MB, D5DA, D5EE, D5ME, D5SX, D5SE, D5IM, DTSX, DTEQ. Procesamiento REDATAM del CNPHV 2022 (INDEC).", "url": "https://celade.cepal.org/bdcelade/MIALC/"},
        {"id": "mialc_2010_dam", "cita": "CEPAL/CELADE. MIALC 2.0. Argentina 2010, DAM: ARG_2010_D5MB, D5DA, D5EE, D5ME, D5SX, DTSX. Procesamiento REDATAM del Censo 2010 (INDEC).", "url": "https://celade.cepal.org/bdcelade/MIALC/"},
        {"id": "mialc_2010_dame", "cita": "CEPAL/CELADE. MIALC 2.0. Argentina 2010, DAME (municipios): ARG_2010_E5MB, matriz básica 2.238 × 2.238.", "url": "https://celade.cepal.org/bdcelade/MIALC/xls/arg/arg2010/ARG_2010_E5MB.xlsX"},
        {"id": "mialc_2001", "cita": "CEPAL/CELADE. MIALC 2.0. Argentina 2001, DAM: ARG_2001_D5MB.", "url": "https://celade.cepal.org/bdcelade/MIALC/"},
        {"id": "mialc_metadatos", "cita": "CEPAL/CELADE (2026). Metadatos MIALC, enero 2026.", "url": "https://celade.cepal.org/bdcelade/MIALC/docs/Metadatos_MIALC_Enero2026.pdf"},
        {"id": "cepal_nota_2022", "cita": "CEPAL (nota informativa). Datos censales de Argentina 2022, Paraguay 2022 y Bolivia 2024 muestran ... (incorporación a MIALC; caída de la intensidad de la migración interna en Argentina 2010-2022). Consultada vía buscador; el sitio bloqueó la descarga directa.", "url": "https://www.cepal.org/es/notas/datos-censales-argentina-2022-paraguay-2022-estado-plurinacional-bolivia-2024-muestran"},
        {"id": "indec_2022_migraciones", "cita": "INDEC (2024). Censo Nacional de Población, Hogares y Viviendas 2022. Resultados definitivos: Migraciones internacionales e internas. Buenos Aires: INDEC. (Copia consultada en sitio de terceros.)", "url": "https://elsubmarinojujuy.com.ar/wp-content/uploads/censo2022_migraciones.pdf"},
        {"id": "renaper_interprov_2023", "cita": "Dirección Nacional de Población – RENAPER (2023). Movilidad residencial interprovincial en Argentina. Septiembre 2023. Base: cambios de domicilio en trámites de DNI, ene-2012 a ago-2022.", "url": "https://www.argentina.gob.ar/sites/default/files/2023/09/movilidad_residencial_interprovincial_en_argentina.pdf"},
        {"id": "renaper_interdep_2023", "cita": "Dirección Nacional de Población – RENAPER (2023). Movilidad residencial interdepartamental en Argentina. Período 2012-2022. Octubre 2023.", "url": "https://www.argentina.gob.ar/sites/default/files/2023/11/movilidad_residencial_interdepartamental_en_argentina.pdf"},
        {"id": "citypop", "cita": "Brinkhoff, T. City Population – Argentina: departamentos y localidades, censos 1991, 2001, 2010 y 2022 (en base a INDEC).", "url": "https://citypopulation.de/en/argentina/admin/"},
        {"id": "tdf_doctec15", "cita": "Gobierno de Tierra del Fuego. Documento técnico 15: Población relativa (series censales 1970-2010).", "url": "https://www.tierradelfuego.gob.ar/wp-content/uploads/2016/08/DocTec-15-Poblaci¦n-relativa.pdf"},
        {"id": "formosa_cuadro_hist", "cita": "Gobierno de Formosa. Cuadro de población por provincia, censos 1947-2010 (en base a INDEC). Usado sólo para Neuquén 1970 y 1980.", "url": "https://archivos.formosa.gob.ar/media/uploads/documentos/documento_1411413239.pdf"},
        {"id": "rodriguez_2004", "cita": "Rodríguez Vignoli, J. (2004). Migración interna en América Latina y el Caribe: estudio regional del período 1980-2000. Serie Población y Desarrollo 50, CEPAL.", "url": "https://repositorio.cepal.org/bitstream/handle/11362/7188/1/S04117_es.pdf"},
        {"id": "copade_vaca_muerta", "cita": "COPADE Neuquén (2024). Ficha Región Vaca Muerta (Censo 2022).", "url": "https://www.copade.gob.ar/wp-content/uploads/2024/05/Ficha-REGION-VACA-MUERTA-.pdf"},
        {"id": "rionegro_anelo", "cita": "Diario Río Negro. 'El Censo en Añelo no reflejó el boom migratorio de Vaca Muerta: «No se condice con la realidad»'.", "url": "https://www.rionegro.com.ar/politica/el-censo-en-anelo-no-reflejo-el-boom-migratorio-de-vaca-muerta-no-se-condice-con-la-realidad-3696938/"},
        {"id": "molinatti_2017", "cita": "Molinatti, F. (2017). Las migraciones internas en Argentina: posibilidades, alcances y desafíos para su captación mediante el Censo 2010. XIV Jornadas Argentinas de Estudios de Población (AEPA), Santa Fe.", "url": "https://www.aacademica.org/xivjornadasaepa/39.pdf"},
        {"id": "perez_2009", "cita": "Pérez, V. del V. (2009). Caracterización de los migrantes intraprovinciales residentes en el Gran San Miguel de Tucumán-Tafí Viejo y el Gran Resistencia, según el Censo 2001. X Jornadas AEPA, Catamarca (resume a Vapñarsky 1994, 1995 y Bertoncello 1994).", "url": "https://www.aacademica.org/000-058/14.pdf"},
        {"id": "gutierrez_rearte_2012", "cita": "Gutiérrez, A. y Rearte, J. F. (2012). Presentación: Dossier Transporte y movilidades en ciudades intermedias de Argentina. Transporte y Territorio 7: 1-8 (cita a Vapñarsky y Gorojovsky 1990; Bertoncello 2011).", "url": "https://www.redalyc.org/pdf/3330/333027353001.pdf"},
        {"id": "quiroga_rosales_2015", "cita": "Quiroga, D. y Rosales (2015). La medición de la migración interna en los censos argentinos 1947-2010. XIII Jornadas AEPA, Salta.", "url": "https://www.aacademica.org/xiiijornadasaepa/36.pdf"},
        {"id": "hermida_2013", "cita": "Hermida, M., Malizia, M. y van Aert, P. (2013) y Hermida (varios). Migración y Ley 19.640 en Tierra del Fuego.", "url": "https://www.aacademica.org/marianohermida/4"},
        {"id": "aletheia_19640", "cita": "Aletheia (UNLP, 2024). Ley 19.640: consecuencias sociales, políticas y económicas; migración interna y déficit habitacional en Tierra del Fuego.", "url": "https://aletheia.fahce.unlp.edu.ar/article/download/ALEe197/19680/58575"},
        {"id": "otero_gonzalez", "cita": "Otero, A. y González, R. (Univ. Nac. del Comahue) y otros. Migración de amenidad en la Norpatagonia andina (San Martín de los Andes, Villa La Angostura, El Bolsón, Bariloche).", "url": "https://cdsa.aacademica.org/000-074/746.pdf"},
        {"id": "kloster_uncoma", "cita": "Revista Geográfica (UNComahue). Crecimiento poblacional de Neuquén 1970-2001 (variaciones intercensales provinciales y peso de Confluencia).", "url": "https://revele.uncoma.edu.ar/index.php/geografia/article/download/237/223/663"},
    ],
}

json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# ---------- stats for the doc
st = {
    "inter22": inter22, "inter10": inter10, "inter10amba": inter10amba,
    "corr22": corridors(mat22, PROV, 20), "corr10amba": corridors(R, order, 20),
    "s22": s22, "s10a": s10a, "masc": masc, "fit22": fit22, "fit10": fit10, "dist": dist_shares, "bands": bands22,
    "orig_shares22": {d: origin_shares(mat22, PROV, d) for d in ["Neuquén", "Tierra del Fuego, Antártida e Islas del Atlántico Sur", "Santa Cruz", "Chubut", "Córdoba", "San Luis", "Río Negro", "Buenos Aires", "Ciudad Autónoma de Buenos Aires", "Mendoza", "Salta"]},
    "orig_shares10a": {d: origin_shares(R, order, d) for d in ["AMBA", "Resto de Buenos Aires", "Neuquén", "Tierra del Fuego, Antártida e Islas del Atlántico Sur", "Santa Cruz", "Córdoba", "Chubut"]},
    "dest_shares10a_AMBA": sorted(((d, R["AMBA"][d]) for d in order if d != "AMBA"), key=lambda x: -x[1])[:10],
    "fam": [fam22, fam10], "typ": [t22, ts22, t10, ts10],
    "ritmos": ritmos,
}
json.dump(st, open(STATS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("ok", inter22, inter10, inter10amba)
